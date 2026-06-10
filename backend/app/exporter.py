import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.repository import Repository


EXPORT_COLUMNS = [
    "id",
    "question",
    "natural_answer",
    "business_intent",
    "tables",
    "sql",
    "answer_summary",
    "result_preview",
    "row_count",
    "status",
    "error",
    "created_at",
]


def export_run_jsonl(repo: Repository, run_id: int, export_dir: Path) -> Path:
    export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / f"run-{run_id}.jsonl"
    questions = repo.list_questions(run_id)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for question in questions:
            handle.write(json.dumps(_export_record(question), ensure_ascii=False, default=str) + "\n")
    return path


def export_run_xlsx(repo: Repository, run_id: int, export_dir: Path) -> Path:
    export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / f"run-{run_id}.xlsx"
    questions = [_export_record(question) for question in repo.list_questions(run_id)]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "benchmark"
    sheet.append(EXPORT_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="111111")

    for question in questions:
        sheet.append(
            [
                json.dumps(question[column], ensure_ascii=False, default=str)
                if isinstance(question.get(column), (dict, list))
                else question.get(column)
                for column in EXPORT_COLUMNS
            ]
        )

    widths = {
        "A": 10,
        "B": 36,
        "C": 52,
        "D": 28,
        "E": 20,
        "F": 58,
        "G": 32,
        "H": 42,
        "I": 12,
        "J": 12,
        "K": 28,
        "L": 22,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    workbook.save(path)
    return path


def _export_record(question: dict) -> dict:
    record = dict(question)
    record["natural_answer"] = record.get("answer_summary") or ""
    return record
